from __future__ import annotations

import json
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(r"D:\HG_Project\etl_pipeline")
ROOT = WORKSPACE / "dwh-pipeline-mapping" / "dwh-pipeline"
WORK = WORKSPACE / "work" / "hg_data_spec"
REFERENCE = Path(r"C:\Users\Admin\Desktop\New folder\SD_26_GIMASYS_01_TÀI_LIỆU_MÔ_TẢ_DỮ_LIỆU_DỰ_ÁN_GMS.docx")
OUTPUT = WORKSPACE / "outputs" / "TAI_LIEU_DAC_TA_DU_LIEU_DU_AN_HG_MEDIA.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
TEXT = "1F1F1F"
RED = "C00000"


TABLE_ALIASES = {
    "dim_po": ["dim_po", "dim_purchase_order", "fact_po"],
    "fact_po_detail": ["fact_po_detail", "fact_purchase_order_detail"],
    "dim_so": ["dim_so", "dim_sale_order", "fact_so"],
    "fact_so_detail": ["fact_so_detail", "fact_sale_order_detail"],
    "dim_khsx": ["dim_khsx", "fact_production_plan", "fact_khsx"],
    "fact_khsx_detail": ["fact_khsx_detail", "fact_production_plan_detail"],
    "bridge_bt_vid": ["bridge_bt_vid", "bridge_editing_vid"],
    "dim_ar": ["dim_ar", "dim_a&r"],
    "dim_subproject_stock": ["dim_subproject_stock", "dim_sub_project_stock"],
    "fact_revenue_by_resources": ["fact_revenue_by_resources", "fact_revenue_by_resources "],
}

CODE_ONLY_COLUMNS = {
    "int_purchase_name_map": [
        ("canonical_name", "text", "NO"),
        ("alias_key", "text", "NO"),
    ],
    "fact_powerbi_unified": [
        ("fact_key", "bigint", "NO"), ("metric_date", "date", "NO"),
        ("grain_type", "text", "NO"), ("metric_name", "text", "NO"),
        ("metric_value", "numeric(38,4)", "NO"), ("resource_id", "text", "YES"),
        ("video_id", "text", "YES"), ("isrc", "text", "YES"),
        ("channel_id", "text", "YES"), ("platform", "text", "YES"),
        ("repository_id", "text", "YES"), ("project_id", "text", "YES"),
        ("sub_project_id", "text", "YES"), ("team", "text", "YES"),
        ("process_step", "text", "YES"), ("status", "text", "YES"),
        ("position_group", "text", "YES"),
    ],
    "mart_powerbi_resource_obt": [
        ("resource_key", "bigint", "NO"), ("resource_id", "text", "NO"),
        ("resource_name", "text", "YES"), ("isrc", "text", "YES"),
        ("resource_origin", "text", "YES"), ("repository_type", "text", "YES"),
        ("repository_id", "text", "YES"), ("repository", "text", "YES"),
        ("sub_project_id", "text", "YES"), ("sub_project_name", "text", "YES"),
        ("project_id", "text", "YES"), ("project", "text", "YES"),
        ("stock_status", "text", "YES"), ("distribution_team", "text", "YES"),
        ("ar_name", "text", "YES"), ("seo_name", "text", "YES"),
        ("stock_stored_date", "date", "YES"), ("acceptance_date", "date", "YES"),
        ("last_distribution_date", "date", "YES"), ("first_used_published_date", "date", "YES"),
        ("last_used_published_date", "date", "YES"), ("first_metric_date", "date", "YES"),
        ("last_metric_date", "date", "YES"), ("resource_age_days", "integer", "YES"),
        ("days_since_last_usage", "integer", "YES"), ("inventory_unused_days", "integer", "YES"),
        ("resource_count", "smallint", "NO"), ("mapped_resource_count", "smallint", "NO"),
        ("used_resource_count", "smallint", "NO"), ("unused_resource_count", "smallint", "NO"),
        ("linked_video_count_all", "bigint", "NO"), ("published_resource_video_uses", "bigint", "NO"),
        ("published_channel_count", "bigint", "NO"), ("uses_position_1", "bigint", "NO"),
        ("uses_position_2_3", "bigint", "NO"), ("uses_position_4_5", "bigint", "NO"),
        ("uses_position_gt_5", "bigint", "NO"), ("lifetime_views", "numeric(38,4)", "NO"),
        ("lifetime_revenue", "numeric(38,4)", "NO"), ("standard_cost", "numeric(38,4)", "NO"),
        ("purchase_cost", "numeric(38,4)", "NO"), ("total_cost", "numeric(38,4)", "NO"),
        ("acceptance_score", "numeric", "YES"),
    ],
}

ENTITY_VI = {
    "company": "công ty", "department": "phòng ban", "employee": "nhân viên",
    "user": "người dùng", "project": "dự án", "sub_project": "dự án con",
    "subproject": "dự án con", "repository": "kho tài nguyên", "resource": "tài nguyên",
    "stock": "HG Stock", "channel": "kênh", "video": "video", "artist": "nghệ sĩ",
    "platform": "nền tảng", "partner": "đối tác", "editing": "phiên biên tập",
    "distribution": "phân phối", "purchase": "thu mua", "sale": "đơn bán/sản xuất",
    "order": "đơn hàng", "po": "đơn mua hàng", "so": "lệnh sản xuất",
    "khsx": "kế hoạch sản xuất", "music": "âm nhạc", "evaluation": "đánh giá",
    "revenue": "doanh thu", "view": "lượt xem", "stream": "lượt nghe",
    "cost": "chi phí", "network": "network/NET", "net": "network/NET",
    "cms": "CMS", "isrc": "mã ISRC", "team": "nhóm", "role": "vai trò",
    "position": "vị trí", "folder": "thư mục", "file": "tệp", "record": "bản ghi",
    "metric": "chỉ số", "status": "trạng thái", "type": "loại", "source": "nguồn",
    "amount": "giá trị", "quantity": "số lượng", "count": "số lượng",
    "date": "ngày", "time": "thời gian", "year": "năm", "month": "tháng",
    "name": "tên", "title": "tiêu đề", "code": "mã", "description": "mô tả",
}


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def aliases(name: str) -> set[str]:
    key = norm(name)
    vals = set(TABLE_ALIASES.get(key, [key]))
    vals.add(key)
    return {norm(x) for x in vals}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "A6A6A6")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def format_cell(cell, header: bool = False, small: bool = False) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", 70), ("left", 90), ("bottom", 70), ("right", 90)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
            run.font.size = Pt(7.5 if small else 8.5)
            run.font.bold = header
            run.font.color.rgb = RGBColor(255, 255, 255) if header else RGBColor.from_string(TEXT)


def add_table(doc, headers: list[str], rows: list[list[object]], widths: list[int], small: bool = False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "TableNormal"
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, text in enumerate(headers):
        hdr.cells[i].text = str(text)
        set_cell_shading(hdr.cells[i], BLUE)
        format_cell(hdr.cells[i], header=True, small=small)
    for ri, values in enumerate(rows):
        row = table.add_row()
        for ci, value in enumerate(values):
            row.cells[ci].text = "" if value is None else str(value)
            if ri % 2:
                set_cell_shading(row.cells[ci], PALE_BLUE)
            format_cell(row.cells[ci], small=small)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def set_font(run, size: float = 10.5, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def ensure_bullet_num(doc) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), "bullet"); lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), "•"); lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc"); lvl_jc.set(qn("w:val"), "left"); lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "260"); p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId"); abs_id.set(qn("w:val"), str(abstract_id)); num.append(abs_id)
    numbering.append(num)
    return num_id


BULLET_NUM_ID = None


def add_bullet(doc, text: str):
    global BULLET_NUM_ID
    if BULLET_NUM_ID is None:
        BULLET_NUM_ID = ensure_bullet_num(doc)
    p = doc.add_paragraph(style="Normal")
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); num_pr.append(ilvl)
    num_id = OxmlElement("w:numId"); num_id.set(qn("w:val"), str(BULLET_NUM_ID)); num_pr.append(num_id)
    p_pr.append(num_pr)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text))
    return p


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        set_font(r, size=14 if level == 1 else (12 if level == 2 else 10.5), bold=True, color=BLUE)
    return p


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_font(run, 9)


def clear_body(doc) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def model_description(name: str) -> str:
    exact = {
        "fact_powerbi_unified": "Bảng fact dạng hẹp hợp nhất các chỉ số cho Power BI; grain_type xác định cấp độ dữ liệu và metric_name xác định chỉ số cộng gộp trong metric_value.",
        "mart_powerbi_resource_obt": "Bảng import Power BI có đúng một dòng cho mỗi tài nguyên HG Stock, kết hợp thuộc tính tài nguyên với chỉ số sử dụng, hiệu suất và chi phí.",
        "int_purchase_name_map": "Model ephemeral chuẩn hóa tên đối tác/kho thu mua và các bí danh để dùng chung cho dữ liệu tài nguyên mua và chi phí mua.",
        "int_stock_video": "Bảng ánh xạ duy nhất giữa tài nguyên HG Stock và video đã liên kết qua mã biên tập.",
        "fact_hao_hut": "Bảng fact phục vụ phân tích hao hụt theo tài nguyên và các bước trong quy trình sản xuất/phân phối.",
        "mart_ton_kho": "Bảng mart theo dõi tài nguyên còn trong kho và tuổi tồn kho.",
        "mart_resource_usage": "Bảng mart tổng hợp mức độ phân phối, sử dụng và hiệu suất của từng tài nguyên.",
        "mart_resource_channel": "Bảng mart tổng hợp quan hệ tài nguyên, video, kênh và các chỉ số hiệu suất.",
        "mart_video_performance_monthly": "Bảng mart tổng hợp lượt xem và doanh thu video theo tháng.",
        "mart_resource_net_monthly": "Bảng mart tổng hợp hiệu suất tài nguyên theo NET và tháng.",
        "mart_repository_usage": "Bảng mart tổng hợp tình trạng sử dụng tài nguyên theo kho.",
        "resources_useage_number": "Bảng tổng hợp số lần sử dụng tài nguyên phục vụ báo cáo.",
    }
    if name in exact:
        return exact[name]
    label = name.replace("_", " ")
    if name.startswith("dim_"):
        return f"Bảng chiều chuẩn hóa thông tin {label[4:]} phục vụ liên kết và phân tích trong DWH HG Media."
    if name.startswith("fact_"):
        return f"Bảng sự kiện ghi nhận dữ liệu nghiệp vụ {label[5:]} ở cấp chi tiết được xác định bởi các khóa trong bảng."
    if name.startswith("mart_"):
        return f"Bảng mart tổng hợp dữ liệu {label[5:]} phục vụ báo cáo và phân tích."
    if name.startswith("bridge_"):
        return f"Bảng cầu nối biểu diễn quan hệ nhiều-nhiều cho {label[7:]}."
    if name.startswith("int_"):
        return f"Model trung gian xử lý {label[4:]} cho các model downstream."
    return f"Object dữ liệu kỹ thuật {name} trong pipeline HG Media."


def infer_column_description(column: str, data_type: str, samples: list[dict]) -> str:
    raw = column.strip()
    key = norm(raw)
    if key in {"_batch_id", "batch_id"}:
        return "Mã lô dữ liệu dùng để truy vết lần extract/load."
    if key in {"_loaded_at", "loaded_at"}:
        return "Thời điểm bản ghi được nạp vào lớp dữ liệu hiện tại."
    if key in {"_source_file", "source_file"}:
        return "Tên tệp nguồn cung cấp bản ghi."
    if key in {"_source_sheet", "source_sheet"}:
        return "Tên sheet nguồn cung cấp bản ghi."
    if key in {"_source_id", "source_id"}:
        return "Mã nguồn dữ liệu trong cấu hình pipeline."
    if key.endswith("_sk"):
        entity = key[:-3].replace("_", " ")
        return f"Khóa thay thế do DWH sinh cho {entity}."
    if key in {"id", "key"} or key.endswith("_id") or key.endswith("id"):
        entity_key = re.sub(r"(_?id|_?key)$", "", key).strip("_") or "bản ghi"
        entity = ENTITY_VI.get(entity_key, entity_key.replace("_", " "))
        sample = next((str(r.get(raw)) for r in samples if r.get(raw) is not None), "")
        kind = "UUID" if re.fullmatch(r"[0-9a-fA-F-]{32,36}", sample) else "định danh"
        return f"Mã {kind} của {entity}."
    if key.startswith("is_") or key.startswith("has_") or "boolean" in data_type.lower():
        return f"Cờ xác định trạng thái {key.replace('_', ' ')}."
    if key.endswith("_at") or "timestamp" in data_type.lower():
        return f"Thời điểm {key.removesuffix('_at').replace('_', ' ')}."
    if key.endswith("_date") or key.startswith("date_") or data_type.lower() == "date":
        return f"Ngày {key.replace('_date', '').replace('date_', '').replace('_', ' ')}."
    if key.endswith("_name") or key == "name" or key.endswith("name"):
        entity_key = re.sub(r"_?name$", "", key).strip("_") or "đối tượng"
        entity = ENTITY_VI.get(entity_key, entity_key.replace("_", " "))
        return f"Tên hiển thị của {entity}."
    if key.endswith("_code") or key == "code":
        entity_key = re.sub(r"_?code$", "", key).strip("_") or "đối tượng"
        entity = ENTITY_VI.get(entity_key, entity_key.replace("_", " "))
        return f"Mã nghiệp vụ của {entity}."
    if "url" in key or key in {"link", "hg_link"}:
        return f"Đường dẫn URL liên quan đến {key.replace('_', ' ')}."
    if "email" in key:
        return "Địa chỉ thư điện tử liên quan đến bản ghi."
    if any(x in key for x in ["revenue", "earning"]):
        return "Giá trị doanh thu ghi nhận cho cấp dữ liệu của bản ghi."
    if "cost" in key or "price" in key or "amount" in key:
        return f"Giá trị tiền tệ của {key.replace('_', ' ')}."
    if "view" in key:
        return "Số lượt xem được ghi nhận cho đối tượng và kỳ dữ liệu tương ứng."
    if "stream" in key:
        return "Số lượt nghe/stream được ghi nhận cho đối tượng và kỳ dữ liệu tương ứng."
    if key.endswith("_count") or key.startswith("number_") or key.startswith("total_"):
        return f"Số lượng/tổng giá trị của {key.replace('_', ' ')}."
    if "status" in key or "state" in key:
        return f"Trạng thái nghiệp vụ của {key.replace('_', ' ')}."
    if key in {"description", "comment", "note", "concluding"} or any(x in key for x in ["description", "comment"]):
        return f"Nội dung mô tả/ghi chú của {key.replace('_', ' ')}."
    if key == "active":
        return "Cờ cho biết bản ghi đang có hiệu lực."
    human = key.replace("_", " ") or raw
    translated = " ".join(ENTITY_VI.get(token, token) for token in human.split())
    return f"Giá trị thuộc tính {translated} của bản ghi; ý nghĩa được suy luận từ schema và dữ liệu mẫu."


def build_lookup(scope: dict):
    catalog_desc = {}
    for row in scope["dictionary"]["catalog"]:
        catalog_desc[norm(row["name"])] = row["description"]
    col_desc = {}
    for row in scope["dictionary"].get("column_catalog", []):
        if row.get("description"):
            col_desc[(norm(row["table"]), norm(row["column"]))] = row["description"]
    detail = {}
    for sheet in scope["dictionary"]["detail_sheets"]:
        table_keys = aliases(sheet["sheet"]) | aliases(sheet["table_name"])
        for field in sheet["fields"]:
            desc = field.get("description_vi") or field.get("description_en")
            rule = field.get("rule") or ""
            for tk in table_keys:
                detail[(tk, norm(field["target_field"]))] = (desc, rule)
    return catalog_desc, col_desc, detail


def object_description(obj: dict, catalog_desc: dict) -> tuple[str, str]:
    for key in aliases(obj["name"]):
        if catalog_desc.get(key):
            return catalog_desc[key], "Data Dictionary"
    if obj["category"] == "source":
        source = obj.get("source_object") or obj.get("source_id")
        return (
            f"Bảng staging lưu dữ liệu thô từ {obj['source_kind']} — object nguồn {source}; "
            "dùng làm đầu vào cho các model dbt và kiểm tra chất lượng sau nạp.",
            "Suy luận kỹ thuật",
        )
    return model_description(obj["name"]), "SQL/dbt + suy luận kỹ thuật"


def column_description(obj: dict, col: dict, col_desc: dict, detail: dict) -> tuple[str, str]:
    ck = norm(col["name"])
    for tk in aliases(obj["name"]):
        if (tk, ck) in col_desc:
            return col_desc[(tk, ck)], "Data Dictionary"
        if (tk, ck) in detail:
            desc, rule = detail[(tk, ck)]
            if desc:
                return desc, "Data Dictionary"
            if rule:
                concise = re.sub(r"\s+", " ", rule).strip()
                if len(concise) <= 260:
                    return f"Trường được xác định theo quy tắc: {concise}", "Mapping Data Dictionary"
    return infer_column_description(col["name"], col.get("data_type", ""), obj.get("samples", [])), "Suy luận kỹ thuật"


def object_columns(obj: dict) -> list[dict]:
    if obj.get("physical"):
        return [
            {"name": c["name"], "data_type": c.get("data_type") or c.get("type") or c.get("udt_name"), "nullable": c.get("nullable", "")}
            for c in obj["physical"]["columns"]
        ]
    return [
        {"name": name, "data_type": dtype, "nullable": nullable}
        for name, dtype, nullable in CODE_ONLY_COLUMNS.get(obj["name"], [])
    ]


def object_status(obj: dict) -> str:
    if obj.get("physical"):
        return "Đã vật lý hóa"
    if obj["name"] == "int_purchase_name_map":
        return "Ephemeral (không tạo bảng)"
    return "Có trong code, chưa vật lý hóa"


def size_text(obj: dict) -> str:
    if not obj.get("physical"):
        return "—"
    size = obj["physical"].get("total_bytes", 0)
    return f"{size / (1024 ** 3):.3f} GB" if size >= 1024 ** 3 else f"{size / (1024 ** 2):.2f} MB"


def rows_text(obj: dict) -> str:
    if not obj.get("physical"):
        return "—"
    estimate = obj["physical"].get("estimated_rows", 0)
    if estimate == 0 and obj["physical"].get("total_bytes", 0) > 1024 * 1024:
        return "Chưa có thống kê"
    return f"~{estimate:,}"


def main() -> None:
    scope = json.loads((WORK / "scope_analysis.json").read_text(encoding="utf-8"))
    evidence = json.loads((WORK / "object_evidence.json").read_text(encoding="utf-8"))
    objects = evidence["objects"]
    catalog_desc, col_desc, detail = build_lookup(scope)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for style_name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)

    # Footer fields.
    footer = section.footer
    for p in footer.paragraphs:
        p._element.getparent().remove(p._element)
    add_page_number(footer.add_paragraph())
    first_footer = section.first_page_footer
    for p in first_footer.paragraphs:
        p._element.getparent().remove(p._element)
    fp = first_footer.add_paragraph("Hà Nội - 2026")
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs:
        set_font(r, 9)

    # Cover.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(92)
    set_font(p.add_run("CÔNG TY TNHH GIẢI PHÁP PHÂN TÍCH DỮ LIỆU INSIGHT DATA"), 14, True, BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run("TÀI LIỆU"), 18, True, TEXT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(74)
    set_font(p.add_run("ĐẶC TẢ DỮ LIỆU DỰ ÁN HG MEDIA"), 18, True, BLUE)
    add_table(doc, ["Thông tin tài liệu", "Thông tin tài liệu"], [
        ["Phiên bản tài liệu", ""], ["Ngày lập tài liệu", ""],
        ["Ngày cập nhật cuối", ""], ["Ngày phê duyệt tài liệu", ""],
        ["Đơn vị ban hành", ""],
    ], [3200, 6880])
    doc.add_page_break()

    add_heading(doc, "BẢNG THEO DÕI HIỆU LỰC VÀ THAY ĐỔI", 1)
    add_table(doc, ["STT", "Ngày", "Phiên bản", "Hình thức*", "Mô tả", "Tác giả"], [
        ["1", "", "", "", "", ""]
    ], [600, 1300, 1200, 1300, 3500, 2180])
    add_body(doc, "(*) A (Add) – thêm mới, M (Modify) – sửa đổi, D (Delete) – xóa")
    add_heading(doc, "MỤC LỤC", 1)
    toc_rows = [
        ["1", "TỔNG QUAN DỰ ÁN"], ["2", "DATA MODEL (MÔ HÌNH DỮ LIỆU)"],
        ["3", "DATA DICTIONARY (TỪ ĐIỂN DỮ LIỆU)"], ["4", "DATA FLOW (LUỒNG DỮ LIỆU)"],
        ["5", "ERROR HANDLER (XỬ LÝ LỖI)"], ["A", "PHỤ LỤC"],
    ]
    add_table(doc, ["Mục", "Nội dung"], toc_rows, [900, 9180])
    doc.add_page_break()

    # 1. Overview.
    add_heading(doc, "1. TỔNG QUAN DỰ ÁN", 1)
    add_heading(doc, "1.1. Mục đích tài liệu", 2)
    add_body(doc, "Tài liệu mô tả cấu trúc dữ liệu, mô hình dữ liệu, luồng xử lý và cơ chế kiểm soát lỗi của dự án xây dựng kho dữ liệu và hệ thống báo cáo thông minh HG Media. Nội dung được đối chiếu từ source code pipeline, SQL/dbt, Data Dictionary và schema PostgreSQL đang triển khai.")
    add_heading(doc, "1.2. Phạm vi", 2)
    add_body(doc, "Phạm vi được chốt theo object đang được khai báo trong code, gồm 64 nguồn/staging và 64 model dbt sản xuất. Không bao gồm model trong thư mục test, các bảng backup hoặc object tồn dư trong DB nhưng không còn được khai báo trong code.")
    add_bullet(doc, "Nguồn dữ liệu: Odoo, HG Stock, hệ thống quản lý kênh, hệ thống biên tập, hệ thống khảo sát/đánh giá, API sale, Elasticsearch và Google Sheet.")
    add_bullet(doc, "Kiến trúc: nguồn → Python EL → MinIO bronze/PostgreSQL staging → dbt silver → dbt gold → Power BI/báo cáo.")
    add_bullet(doc, "Snapshot schema và kích thước bảng được lấy ngày 08/09/2026. Số dòng dùng thống kê PostgreSQL; giá trị “Chưa có thống kê” không đồng nghĩa bảng rỗng.")
    add_heading(doc, "1.3. Đối tượng sử dụng", 2)
    add_table(doc, ["Đối tượng", "Mục đích sử dụng"], [
        ["Business Analyst", "Đối chiếu phạm vi, thuật ngữ và logic dữ liệu."],
        ["Data Engineer / Developer", "Phát triển, vận hành và bảo trì pipeline EL/dbt."],
        ["QA / Tester", "Thiết kế kiểm thử schema, dữ liệu và dependency."],
        ["BI Developer / Data Analyst", "Khai thác silver/gold cho báo cáo và phân tích."],
    ], [2800, 7280])

    # 2. Data model.
    add_heading(doc, "2. DATA MODEL (MÔ HÌNH DỮ LIỆU)", 1)
    add_heading(doc, "2.1. Sơ đồ tổng thể", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("DB / API / Elasticsearch / Google Sheet  →  Python EL  →  MinIO Bronze  →  PostgreSQL Staging  →  dbt Silver  →  dbt Gold  →  Power BI"), 10, True, BLUE)
    source_counts = Counter(o["source_kind"] for o in objects if o["category"] == "source")
    model_counts = Counter(o["layer"] for o in objects if o["category"] == "model")
    add_table(doc, ["Nhóm object", "Số lượng", "Vai trò"], [
        ["Database source / staging", source_counts["Database"], "Bảng dữ liệu thô từ các CSDL nghiệp vụ."],
        ["Google Sheet source / staging", source_counts["Google Sheet"], "Dữ liệu nghiệp vụ bổ sung và lịch sử."],
        ["Elasticsearch source / staging", source_counts["Elasticsearch"], "Thông tin và metric video/kênh."],
        ["API source / staging", source_counts["API"], "Dữ liệu sale-details theo kỳ báo cáo."],
        ["dbt silver", model_counts["silver"], "Chuẩn hóa dimension, fact và bridge."],
        ["dbt intermediate", model_counts["intermediate"], "Logic trung gian/ephemeral."],
        ["dbt gold", model_counts["gold"], "Mart và bảng phục vụ báo cáo."],
        ["TỔNG", len(objects), "Toàn bộ object trong phạm vi code."],
    ], [3300, 1200, 5580])

    add_heading(doc, "2.2. Danh sách Object", 2)
    for group_title, predicate in [
        ("Nguồn và staging", lambda o: o["category"] == "source"),
        ("Model silver", lambda o: o["category"] == "model" and o["layer"] == "silver"),
        ("Model intermediate", lambda o: o["category"] == "model" and o["layer"] == "intermediate"),
        ("Model gold", lambda o: o["category"] == "model" and o["layer"] == "gold"),
    ]:
        add_heading(doc, group_title, 3)
        group = [o for o in objects if predicate(o)]
        rows = []
        for i, obj in enumerate(group, 1):
            source_or_refs = obj.get("source_object") or ", ".join(obj.get("refs", [])) or "—"
            target = f"{obj.get('target_schema', obj['layer'])}.{obj['name']}"
            rows.append([i, obj["name"], obj["source_kind"], source_or_refs, target, obj.get("load_mode") or "—", object_status(obj), size_text(obj)])
        add_table(doc, ["STT", "Object", "Loại", "Nguồn / ref", "Đích", "Cách nạp", "Trạng thái", "Kích thước"], rows,
                  [500, 1700, 1100, 2100, 1500, 1000, 1300, 880], small=True)

    add_heading(doc, "2.3. Quan hệ và dependency", 2)
    add_body(doc, "Dependency model được xác định trực tiếp từ hàm ref() trong SQL dbt. dbt xây dựng đồ thị có hướng và chạy model theo thứ tự phụ thuộc; các model không phụ thuộc nhau có thể chạy song song theo cấu hình 4 threads.")
    dep_rows = []
    for obj in [o for o in objects if o["category"] == "model"]:
        dep_rows.append([obj["name"], ", ".join(obj.get("refs", [])) or "Nguồn staging/source()", f"{obj.get('target_schema', obj['layer'])}.{obj['name']}"])
    add_table(doc, ["Model", "Input trực tiếp", "Output"], dep_rows, [2600, 5000, 2480], small=True)

    # 3. Data dictionary.
    add_heading(doc, "3. DATA DICTIONARY (TỪ ĐIỂN DỮ LIỆU)", 1)
    add_heading(doc, "3.1. Hướng dẫn đọc", 2)
    add_body(doc, "Mỗi object có một dòng metadata kỹ thuật và một bảng cột. Tên kỹ thuật lấy từ code/DB; tên trong Excel được dùng làm alias khi có xung đột. Cột “Cơ sở mô tả” phân biệt nội dung có sẵn trong Data Dictionary với nội dung được suy luận kỹ thuật từ schema, dữ liệu mẫu, SQL/dbt hoặc quy tắc mapping.")
    add_body(doc, "Đối với object đã vật lý hóa, kiểu dữ liệu và nullable lấy từ information_schema của PostgreSQL. Đối với object code-only, cấu trúc được lấy từ SELECT cuối của SQL/model YAML khi xác định được.")
    add_heading(doc, "3.2. Quy ước kiểu dữ liệu PostgreSQL", 2)
    add_table(doc, ["Nhóm kiểu", "Ví dụ", "Ý nghĩa"], [
        ["Chuỗi", "text, character varying", "Mã, tên, mô tả và nội dung văn bản."],
        ["Số nguyên", "smallint, integer, bigint", "Số lượng, khóa số hoặc chỉ số nguyên."],
        ["Số thập phân", "numeric, double precision", "Doanh thu, chi phí, tỷ lệ hoặc chỉ số đo lường."],
        ["Ngày/giờ", "date, timestamp", "Ngày nghiệp vụ hoặc thời điểm sự kiện/nạp dữ liệu."],
        ["Boolean", "boolean", "Cờ đúng/sai hoặc trạng thái có/không."],
        ["Cấu trúc", "json/jsonb, ARRAY", "Dữ liệu bán cấu trúc hoặc danh sách giá trị."],
    ], [1800, 2800, 5480])

    add_heading(doc, "3.3. Chi tiết Object và cột dữ liệu", 2)
    for index, obj in enumerate(objects, 1):
        desc, desc_source = object_description(obj, catalog_desc)
        add_heading(doc, f"3.3.{index}. {obj['name']}", 3)
        target = f"{obj.get('target_schema', obj['layer'])}.{obj['name']}"
        meta = (
            f"Loại: {obj['source_kind']} | Đích: {target} | Trạng thái: {object_status(obj)} | "
            f"Số dòng: {rows_text(obj)} | Kích thước: {size_text(obj)} | Mô tả bảng: {desc} "
            f"(Cơ sở: {desc_source})."
        )
        add_body(doc, meta)
        cols = object_columns(obj)
        if not cols:
            add_body(doc, "Chưa xác định được danh sách cột do model chưa vật lý hóa và SELECT cuối không khai báo trực tiếp schema output.")
            continue
        rows = []
        for ci, col in enumerate(cols, 1):
            cdesc, basis = column_description(obj, col, col_desc, detail)
            nullable = col.get("nullable", "")
            nullable_vi = "Có" if str(nullable).upper() in {"YES", "Y", "TRUE"} else ("Không" if str(nullable).upper() in {"NO", "N", "FALSE"} else str(nullable))
            rows.append([ci, col["name"], col.get("data_type", ""), nullable_vi, cdesc, basis])
        add_table(doc, ["STT", "Tên cột", "Kiểu dữ liệu", "Nullable", "Ý nghĩa", "Cơ sở mô tả"], rows,
                  [500, 1800, 1400, 850, 4100, 1430], small=True)

    # 4. Flow.
    add_heading(doc, "4. DATA FLOW (LUỒNG DỮ LIỆU)", 1)
    add_heading(doc, "4.1. Kiến trúc tổng thể", 2)
    add_table(doc, ["Bước", "Thành phần", "Mô tả", "Output"], [
        ["1", "Extract", "Extractor đọc dữ liệu theo cấu hình nguồn; ghi nhận checksum/watermark và batch_id.", "DataFrame hoặc dữ liệu stream"],
        ["2", "Bronze", "Dữ liệu batch được lưu dạng Parquet trên MinIO khi không stream trực tiếp.", "bronze/<type>/<source>/<batch>"],
        ["3", "Load staging", "Nạp dữ liệu thô vào PostgreSQL staging theo truncate hoặc append/upsert đã cấu hình.", "staging.*"],
        ["4", "EL Data Quality", "Kiểm tra schema, batch metadata, khóa, null, uniqueness và miền giá trị theo YAML.", "meta.dq_*"],
        ["5", "dbt transform", "Build các model theo ref() từ silver đến gold.", "silver.*, gold.*"],
        ["6", "dbt/Data Quality", "Chạy dbt test và kiểm tra DQ trên các bảng silver/gold được khai báo.", "Kết quả DQ và trạng thái DAG"],
        ["7", "Consumption", "Power BI và các báo cáo đọc dữ liệu từ lớp gold hoặc bảng tối ưu chuyên dụng.", "Dashboard / báo cáo"],
    ], [700, 1700, 5400, 2280])
    add_heading(doc, "4.2. Lịch chạy", 2)
    add_body(doc, "Các biểu thức dưới đây là cron đúng theo code. Repo không khai báo timezone riêng trong từng DAG; vì vậy tài liệu không quy đổi sang múi giờ cụ thể.")
    add_table(doc, ["DAG", "Cron", "Diễn giải", "Retry"], [
        ["el_api_pipeline", "0 5 * * *", "Hằng ngày lúc 05:00 theo timezone Airflow", "2 lần, cách 10 phút"],
        ["el_google_sheet_pipeline", "0 6 * * *", "Hằng ngày lúc 06:00 theo timezone Airflow", "2 lần, cách 5 phút"],
        ["el_database_pipeline", "0 */4 * * *", "Mỗi 4 giờ", "2 lần, cách 5 phút"],
        ["el_elastic_pipeline", "0 */6 * * *", "Mỗi 6 giờ", "2 lần, cách 10 phút"],
        ["el_csv_pipeline", "0 5 * * *", "Hằng ngày lúc 05:00; không có CSV object trong phạm vi cấu hình được yêu cầu", "Theo code DAG"],
        ["dbt_transform_pipeline", "30 7 * * *", "Hằng ngày lúc 07:30; chờ 4 EL DAG", "1 lần, cách 10 phút"],
        ["data_quality_el_*", "Triggered", "Được EL DAG tương ứng kích hoạt sau load", "Theo DAG DQ"],
        ["data_quality_dbt_pipeline", "Triggered", "Được dbt DAG kích hoạt sau run/test", "Theo DAG DQ"],
    ], [2600, 1400, 4600, 1480])
    add_heading(doc, "4.3. Dependency Map", 2)
    add_body(doc, "EL API / Database / Elasticsearch / Google Sheet → DQ staging tương ứng → dbt_transform_pipeline → dbt test (nếu bật) → data_quality_dbt_pipeline → Power BI/báo cáo.")
    add_body(doc, "dbt_transform_pipeline dùng ExternalTaskSensor: timeout 3.600 giây cho Google Sheet, Database và API; timeout 7.200 giây cho Elasticsearch. Khi chạy thủ công có thể chọn layer/model và bật hoặc tắt bước chờ EL, dbt test.")

    # 5. Error handling.
    add_heading(doc, "5. ERROR HANDLER (XỬ LÝ LỖI)", 1)
    add_heading(doc, "5.1. Phân loại lỗi", 2)
    add_table(doc, ["Mã", "Nhóm lỗi", "Bằng chứng trong code", "Hành vi"], [
        ["ERR-001", "Kết nối/Extract", "Extractor hoặc connector phát sinh exception", "Task fail và Airflow retry theo DAG."],
        ["ERR-002", "Load staging", "run_load bắt exception, ghi log lỗi và raise", "Task fail; không đánh dấu batch load thành công."],
        ["ERR-003", "Schema/metadata batch", "DQ kiểm tra cột, _batch_id, _loaded_at và target table", "Critical làm DQ DAG fail; warning được ghi nhận."],
        ["ERR-004", "Khóa/null/miền giá trị", "Great Expectations từ YAML DQ", "Lưu kết quả từng rule vào meta.dq_validation_results."],
        ["ERR-005", "dbt run/test", "BashOperator chạy dbt; non-zero exit", "dbt task fail và retry một lần."],
        ["ERR-006", "Dependency timeout", "ExternalTaskSensor hết timeout", "dbt workflow không tiếp tục theo nhánh chờ EL."],
        ["ERR-007", "Không có thay đổi/0 dòng", "run_extract kiểm tra checksum/watermark/row_count", "Skip load có kiểm soát; không coi là lỗi kỹ thuật."],
    ], [850, 1800, 4450, 2980])
    add_heading(doc, "5.2. Quy trình xử lý lỗi", 2)
    for text in [
        "Airflow ghi nhận trạng thái task và thực hiện retry theo cấu hình của DAG.",
        "SourceRegistry lưu batch_id, minio_path, row_count, checksum, watermark và trạng thái để truy vết batch.",
        "Kết quả DQ được lưu theo run và rule; mẫu bản ghi không đạt được lưu riêng trong meta.dq_unexpected_samples.",
        "Rule severity=critical có thể làm DAG DQ thất bại; severity=warning được tổng hợp nhưng không tự động dừng pipeline.",
        "Khi cần khôi phục staging, rollback_task tìm batch MinIO hợp lệ theo ngày/batch rồi nạp lại bảng đích.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "5.3. Checklist kiểm tra sau mỗi lần chạy", 2)
    for text in [
        "Các task Extract và Load kết thúc ở trạng thái thành công hoặc skip có lý do.",
        "Batch metadata có batch_id, row_count, checksum/watermark và đường dẫn MinIO khi áp dụng.",
        "Không có critical failure trong DQ staging và DQ silver/gold.",
        "dbt run hoàn tất đúng dependency; dbt test pass đối với các model được chọn.",
        "Số dòng và kích thước các bảng trọng yếu không biến động bất thường so với kỳ trước.",
        "Các model code-only cần được phân biệt rõ với bảng đã vật lý hóa trước khi bàn giao cho BI.",
    ]:
        add_bullet(doc, text)

    # Appendix.
    add_heading(doc, "PHỤ LỤC", 1)
    add_heading(doc, "A. Thuật ngữ viết tắt", 2)
    add_table(doc, ["Thuật ngữ", "Giải thích"], [
        ["DWH", "Data Warehouse – kho dữ liệu tập trung."],
        ["EL", "Extract–Load – trích xuất và nạp dữ liệu thô."],
        ["ETL", "Extract–Transform–Load."],
        ["dbt", "Công cụ transform và quản lý dependency model bằng SQL/ref()."],
        ["Staging", "Lớp lưu dữ liệu thô sau khi nạp từ nguồn."],
        ["Silver", "Lớp dữ liệu chuẩn hóa dimension/fact/bridge."],
        ["Gold", "Lớp mart phục vụ báo cáo và phân tích."],
        ["DQ", "Data Quality – kiểm tra chất lượng dữ liệu."],
        ["ISRC", "International Standard Recording Code – mã định danh bản ghi âm."],
        ["CID", "Content ID – cơ chế nhận diện và quản lý bản quyền nội dung."],
        ["NET", "Mạng lưới/CMS quản lý các kênh YouTube."],
        ["DSP", "Digital Service Provider – nền tảng phân phối nội dung số."],
        ["OBT", "One Big Table – bảng rộng tối ưu cho công cụ BI."],
    ], [2100, 7980])
    add_heading(doc, "B. Object chưa vật lý hóa", 2)
    missing = [o for o in objects if not o.get("physical")]
    add_table(doc, ["Object", "Loại", "Trạng thái", "Đường dẫn code"], [
        [o["name"], o["layer"], object_status(o), o.get("path", "")]
        for o in missing
    ], [2400, 1400, 2600, 3680])
    add_heading(doc, "C. Nguồn tài liệu", 2)
    for source in [
        str(REFERENCE),
        r"C:\Users\Admin\Desktop\New folder\TÀI LIỆU GIỚI THIỆU DỰ ÁN HG MEDIA.docx",
        r"C:\Users\Admin\Desktop\New folder\Data Dictionary.xlsx",
        str(ROOT / "config" / "api_sources.yaml"),
        str(ROOT / "config" / "db_sources.yaml"),
        str(ROOT / "config" / "elastic_sources.yaml"),
        str(ROOT / "config" / "google_sheet_sources.yaml"),
        str(ROOT / "src" / "dags"),
        str(ROOT / "dwh_dbt" / "models"),
        "PostgreSQL information_schema và pg_stat_user_tables (snapshot 08/09/2026).",
    ]:
        add_bullet(doc, source)

    # Final document defaults and field refresh.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.core_properties.title = "Tài liệu đặc tả dữ liệu dự án HG Media"
    doc.core_properties.subject = "Data Model, Data Dictionary, Data Flow và Error Handler"
    doc.save(OUTPUT)
    print(json.dumps({
        "output": str(OUTPUT), "objects": len(objects),
        "columns": sum(len(object_columns(o)) for o in objects),
        "tables": len(doc.tables), "paragraphs": len(doc.paragraphs),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
