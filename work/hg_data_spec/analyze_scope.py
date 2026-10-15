from __future__ import annotations

import json
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\HG_Project\etl_pipeline\dwh-pipeline-mapping\dwh-pipeline")
WORK = Path(r"D:\HG_Project\etl_pipeline\work\hg_data_spec")
DICT = Path(r"C:\Users\Admin\Desktop\New folder\Data Dictionary.xlsx")


def norm(value: object) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def yaml_sources() -> dict:
    groups = {}
    for name, key in [
        ("api_sources.yaml", "api_sources"),
        ("db_sources.yaml", "db_sources"),
        ("elastic_sources.yaml", "elastic_sources"),
        ("google_sheet_sources.yaml", "google_sheet_sources"),
    ]:
        rows = []
        current = None
        for raw in (ROOT / "config" / name).read_text(encoding="utf-8").splitlines():
            line = raw.split("#", 1)[0].rstrip()
            match_item = re.match(r"\s*-\s+source_id:\s*(.+?)\s*$", line)
            if match_item:
                if current:
                    rows.append(current)
                current = {"source_id": match_item.group(1).strip().strip('"').strip("'")}
                continue
            if current:
                match_kv = re.match(r"\s+([A-Za-z_][A-Za-z0-9_]*):\s*(.*?)\s*$", line)
                if match_kv:
                    k, v = match_kv.groups()
                    v = v.strip().strip('"').strip("'")
                    if v in {"null", "~", ""}:
                        current[k] = None
                    elif v.lower() in {"true", "false"}:
                        current[k] = v.lower() == "true"
                    else:
                        current[k] = v
        if current:
            rows.append(current)
        groups[key] = [
            {
                "source_id": x.get("source_id"),
                "source_type": x.get("source_type"),
                "connection": x.get("connection"),
                "source_table": x.get("source_table") or x.get("es_index") or x.get("worksheet_name") or x.get("worksheet_pattern") or x.get("file_name") or x.get("api_url"),
                "target_staging_table": x.get("target_staging_table"),
                "load_mode": x.get("load_mode"),
                "watermark_column": x.get("watermark_column"),
                "incremental": x.get("incremental"),
            }
            for x in rows
        ]
    return groups


def model_inventory() -> list[dict]:
    out = []
    model_root = ROOT / "dwh_dbt" / "models"
    for path in sorted(model_root.rglob("*.sql")):
        text = path.read_text(encoding="utf-8", errors="replace")
        refs = sorted(set(re.findall(r"ref\(\s*['\"]([^'\"]+)", text)))
        sources = sorted(set(".".join(x) for x in re.findall(r"source\(\s*['\"]([^'\"]+)['\"]\s*,\s*['\"]([^'\"]+)", text)))
        materialized = re.search(r"materialized\s*=\s*['\"]([^'\"]+)", text)
        out.append(
            {
                "name": path.stem,
                "layer": path.relative_to(model_root).parts[0],
                "path": str(path.relative_to(ROOT)),
                "refs": refs,
                "sources": sources,
                "materialized": materialized.group(1) if materialized else None,
                "has_todo": bool(re.search(r"TODO|skeleton", text, re.I)),
                "select_star": bool(re.search(r"select\s+\*", text, re.I)),
                "line_count": text.count("\n") + 1,
            }
        )
    return out


def dictionary_inventory() -> dict:
    wb = load_workbook(DICT, read_only=False, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        if ws.title in {"Tables", "Columns"}:
            continue
        table_name = norm(ws.cell(1, 2).value) or ws.title
        table_meaning = norm(ws.cell(2, 2).value)
        headers = [norm(ws.cell(3, c).value) for c in range(1, ws.max_column + 1)]
        field_rows = []
        for r in range(4, ws.max_row + 1):
            vals = [norm(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 12) + 1)]
            target_field = vals[2] if len(vals) > 2 else ""
            if not target_field:
                continue
            field_rows.append(
                {
                    "target_field": target_field,
                    "description_en": vals[3] if len(vals) > 3 else "",
                    "description_vi": vals[4] if len(vals) > 4 else "",
                    "data_type": vals[5] if len(vals) > 5 else "",
                    "source_system": vals[6] if len(vals) > 6 else "",
                    "source_table": vals[7] if len(vals) > 7 else "",
                    "source_field": vals[8] if len(vals) > 8 else "",
                    "rule": vals[10] if len(vals) > 10 else "",
                }
            )
        sheets.append(
            {
                "sheet": ws.title,
                "table_name": table_name,
                "table_meaning": table_meaning,
                "headers": headers,
                "field_count": len(field_rows),
                "missing_description_en": sum(not r["description_en"] for r in field_rows),
                "missing_description_vi": sum(not r["description_vi"] for r in field_rows),
                "missing_data_type": sum(not r["data_type"] for r in field_rows),
                "missing_source_system": sum(not r["source_system"] for r in field_rows),
                "missing_source_table": sum(not r["source_table"] for r in field_rows),
                "missing_source_field": sum(not r["source_field"] for r in field_rows),
                "missing_rule": sum(not r["rule"] for r in field_rows),
                "fields": field_rows,
            }
        )

    ws_tables = wb["Tables"]
    catalog = []
    for row in ws_tables.iter_rows(min_row=2, values_only=True):
        if norm(row[1] if len(row) > 1 else None):
            catalog.append(
                {
                    "name": norm(row[1]),
                    "description": norm(row[2] if len(row) > 2 else None),
                    "type": norm(row[3] if len(row) > 3 else None),
                }
            )
    ws_columns = wb["Columns"]
    column_catalog = []
    for row in ws_columns.iter_rows(min_row=2, values_only=True):
        table_name = norm(row[1] if len(row) > 1 else None)
        column_name = norm(row[2] if len(row) > 2 else None)
        if table_name and column_name:
            column_catalog.append(
                {
                    "table": table_name,
                    "column": column_name,
                    "data_type": norm(row[3] if len(row) > 3 else None),
                    "description": norm(row[4] if len(row) > 4 else None),
                    "nullable": norm(row[5] if len(row) > 5 else None),
                    "data_length": norm(row[6] if len(row) > 6 else None),
                }
            )
    wb.close()
    return {"catalog": catalog, "column_catalog": column_catalog, "detail_sheets": sheets}


def read_env(path: Path) -> dict[str, str]:
    env = {}
    if not path.exists():
        return env
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip().strip('"').strip("'")
    return env


def postgres_inventory() -> dict:
    env = read_env(ROOT / ".env")
    required = ["DWH_PG_HOST", "DWH_PG_PORT", "DWH_PG_DB", "DWH_PG_USER", "DWH_PG_PASSWORD"]
    if any(not env.get(k) for k in required):
        return {"status": "missing_env", "missing": [k for k in required if not env.get(k)]}
    try:
        import psycopg2

        conn = psycopg2.connect(
            host=env["DWH_PG_HOST"],
            port=env["DWH_PG_PORT"],
            dbname=env["DWH_PG_DB"],
            user=env["DWH_PG_USER"],
            password=env["DWH_PG_PASSWORD"],
            connect_timeout=8,
        )
        conn.set_session(readonly=True, autocommit=True)
        cur = conn.cursor()
        cur.execute(
            """
            select table_schema, table_name
            from information_schema.tables
            where table_schema in ('staging','silver','gold','meta')
              and table_type='BASE TABLE'
            order by table_schema, table_name
            """
        )
        tables = []
        for schema, table in cur.fetchall():
            cur.execute(
                """
                select column_name, data_type, is_nullable
                from information_schema.columns
                where table_schema=%s and table_name=%s
                order by ordinal_position
                """,
                (schema, table),
            )
            columns = [{"name": a, "type": b, "nullable": c} for a, b, c in cur.fetchall()]
            cur.execute('select count(*) from "{}"."{}"'.format(schema.replace('"','""'), table.replace('"','""')))
            row_count = cur.fetchone()[0]
            tables.append({"schema": schema, "table": table, "row_count": row_count, "columns": columns})
        cur.close()
        conn.close()
        return {"status": "ok", "tables": tables}
    except Exception as exc:
        return {"status": "error", "error_type": type(exc).__name__, "message": str(exc)}


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    result = {
        "sources": yaml_sources(),
        "models": model_inventory(),
        "dictionary": dictionary_inventory(),
        "postgres": postgres_inventory(),
    }
    (WORK / "scope_analysis.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "source_counts": {k: len(v) for k, v in result["sources"].items()},
        "model_counts": dict(Counter(x["layer"] for x in result["models"])),
        "model_total": len(result["models"]),
        "todo_models": sum(x["has_todo"] for x in result["models"]),
        "dictionary_catalog": len(result["dictionary"]["catalog"]),
        "dictionary_detail_sheets": len(result["dictionary"]["detail_sheets"]),
        "postgres_status": result["postgres"].get("status"),
        "postgres_tables": len(result["postgres"].get("tables", [])),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
