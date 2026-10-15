from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(r"D:\HG_Project\etl_pipeline")
OUT = ROOT / "work" / "hg_data_spec"

DOCS = {
    "gms_template": Path(r"C:\Users\Admin\Desktop\New folder\SD_26_GIMASYS_01_TÀI_LIỆU_MÔ_TẢ_DỮ_LIỆU_DỰ_ÁN_GMS.docx"),
    "hg_intro": Path(r"C:\Users\Admin\Desktop\New folder\TÀI LIỆU GIỚI THIỆU DỰ ÁN HG MEDIA.docx"),
}
XLSX = Path(r"C:\Users\Admin\Desktop\New folder\Data Dictionary.xlsx")


def clean(text: object) -> str:
    return re.sub(r"\s+", " ", "" if text is None else str(text)).strip()


def docx_summary(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        text = clean(p.text)
        if text:
            paragraphs.append({"index": i, "style": p.style.name if p.style else None, "text": text})

    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        tables.append({"index": ti, "rows": rows})

    sections = []
    for si, sec in enumerate(doc.sections):
        sections.append(
            {
                "index": si,
                "orientation": str(sec.orientation),
                "page_width": sec.page_width,
                "page_height": sec.page_height,
                "top_margin": sec.top_margin,
                "bottom_margin": sec.bottom_margin,
                "left_margin": sec.left_margin,
                "right_margin": sec.right_margin,
                "header_distance": sec.header_distance,
                "footer_distance": sec.footer_distance,
                "header": [clean(p.text) for p in sec.header.paragraphs if clean(p.text)],
                "footer": [clean(p.text) for p in sec.footer.paragraphs if clean(p.text)],
            }
        )

    props = doc.core_properties
    return {
        "path": str(path),
        "core_properties": {
            "title": props.title,
            "subject": props.subject,
            "author": props.author,
            "keywords": props.keywords,
            "last_modified_by": props.last_modified_by,
            "created": str(props.created),
            "modified": str(props.modified),
        },
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": sections,
        "inline_shapes": len(doc.inline_shapes),
    }


def workbook_summary(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=False)
    result = {"path": str(path), "sheets": []}
    for ws in wb.worksheets:
        samples = []
        nonempty = 0
        for row in ws.iter_rows(values_only=True):
            vals = [clean(v) for v in row]
            if any(vals):
                nonempty += 1
                if len(samples) < 25:
                    samples.append(vals[:40])
        result["sheets"].append(
            {
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_rows": nonempty,
                "samples": samples,
            }
        )
    wb.close()
    return result


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, path in DOCS.items():
        (OUT / f"{name}.json").write_text(
            json.dumps(docx_summary(path), ensure_ascii=False, indent=2), encoding="utf-8"
        )
    (OUT / "data_dictionary.json").write_text(
        json.dumps(workbook_summary(XLSX), ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
