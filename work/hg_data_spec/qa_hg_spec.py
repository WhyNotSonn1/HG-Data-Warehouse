from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


WORK = Path(r"D:\HG_Project\etl_pipeline\work\hg_data_spec")
DOCX = Path(r"D:\HG_Project\etl_pipeline\outputs\TAI_LIEU_DAC_TA_DU_LIEU_DU_AN_HG_MEDIA.docx")
REFERENCE = Path(r"C:\Users\Admin\Desktop\New folder\SD_26_GIMASYS_01_TÀI_LIỆU_MÔ_TẢ_DỮ_LIỆU_DỰ_ÁN_GMS.docx")


def main() -> None:
    doc = Document(DOCX)
    evidence = json.loads((WORK / "object_evidence.json").read_text(encoding="utf-8"))["objects"]
    headings = [p.text.strip() for p in doc.paragraphs if p.style and p.style.name == "Heading 3"]
    dictionary_headings = [x for x in headings if re.match(r"^3\.3\.\d+\. ", x)]
    names = [re.sub(r"^3\.3\.\d+\. ", "", x) for x in dictionary_headings]
    expected = [o["name"] for o in evidence]
    missing = [name for name in expected if name not in names]
    unexpected = [name for name in names if name not in expected]
    duplicates = sorted({name for name in names if names.count(name) > 1})
    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            all_text += "\n" + " | ".join(cell.text for cell in row.cells)
    forbidden = [x for x in ["row_count.gmail", "row reminders", "DỰ ÁN GMS", "BigQuery → ClickHouse"] if x in all_text]
    blank_metadata = []
    for row in doc.tables[0].rows[1:]:
        blank_metadata.append(row.cells[1].text == "")
    expected_columns = 0
    for o in evidence:
        if o.get("physical"):
            expected_columns += len(o["physical"]["columns"])
        else:
            expected_columns += {"int_purchase_name_map": 2, "fact_powerbi_unified": 17, "mart_powerbi_resource_obt": 43}.get(o["name"], 0)
    # Dictionary tables are the last 128 object-specific tables before section 4; identify by header labels.
    dictionary_tables = [
        t for t in doc.tables
        if len(t.rows) and [c.text for c in t.rows[0].cells][:2] == ["STT", "Tên cột"]
    ]
    documented_columns = sum(max(0, len(t.rows) - 1) for t in dictionary_tables)
    ref_hash = hashlib.sha256(REFERENCE.read_bytes()).hexdigest().upper()
    result = {
        "docx_exists": DOCX.exists(),
        "docx_bytes": DOCX.stat().st_size,
        "sections": len(doc.sections),
        "tables": len(doc.tables),
        "paragraphs": len(doc.paragraphs),
        "dictionary_headings": len(dictionary_headings),
        "expected_objects": len(expected),
        "missing_objects": missing,
        "unexpected_objects": unexpected,
        "duplicate_objects": duplicates,
        "dictionary_tables": len(dictionary_tables),
        "expected_columns": expected_columns,
        "documented_columns": documented_columns,
        "metadata_blank": all(blank_metadata),
        "forbidden_text": forbidden,
        "reference_hash_unchanged": ref_hash == "9E3E12EFA41DF3E34C59EB18FB0FF0099A42FFD3245C1FB2D770614C3D010678",
        "update_fields": bool(doc.settings._element.find(qn("w:updateFields")) is not None),
    }
    result["passed"] = all([
        not missing, not unexpected, not duplicates, len(dictionary_headings) == len(expected),
        len(dictionary_tables) == len(expected), documented_columns == expected_columns,
        result["metadata_blank"], not forbidden, result["reference_hash_unchanged"], result["update_fields"],
    ])
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
